import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

import os
# Veri şablonlarını (Excel) üretir: python3 scripts/build-data-templates.py
ROOT=os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
provinces=json.load(open(f'{ROOT}/packages/shared/data/locations/provinces.json'))
districts=json.load(open(f'{ROOT}/packages/shared/data/locations/districts.json'))
pname={p['code']:p['name'] for p in provinces}

F='Arial'
H=Font(name=F,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='0B3B47')
REQ=PatternFill('solid',fgColor='FFF2CC')
EX=Font(name=F,italic=True,color='7F7F7F')
N=Font(name=F); B=Font(name=F,bold=True)
thin=Side(style='thin',color='D9D9D9'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)

wb=Workbook()
ws=wb.active; ws.title='Okuyun'

COMMON_TAIL=[('kaynak',True,'Bilginin alındığı sayfanın tam adresi (https://...). Resmî site öncelikli. Kaynağı olmayan satır kabul edilmez.'),
             ('kaynak_tarihi',True,'Kaynağa bakılan gün: GG.AA.YYYY (örnek 27.10.2026).'),
             ('not',False,'İsteğe bağlı kısa not (örnek: "telefon teyit edilmeli").')]
LOC=[('il',True,'İl adı, "İlçeler" sekmesindeki yazımla (örnek: Bursa).'),
     ('ilce',True,'İlçe adı, "İlçeler" sekmesindeki yazımla (örnek: Nilüfer). Merkez ilçeli illerde "Merkez".'),
     ('mahalle',False,'Mahalle adı, biliniyorsa (örnek: Ataevler).'),
     ('adres',False,'Açık adres.')]
GEO=[('enlem',False,'Harita konumu, ondalık derece (örnek: 40.1885). Bilinmiyorsa boş bırakın, uydurmayın.'),
     ('boylam',False,'Harita konumu, ondalık derece (örnek: 29.0610).')]
CONTACT=[('telefon',False,'Yalnızca kurumun resmî numarası, rakamlarla (örnek: 02242251900). Kişisel numara yazılmaz.'),
         ('web',False,'Resmî web adresi.')]

SHEETS=[
 ('Önemli Kurumlar','Valilik, kaymakamlık, il müdürlükleri, hastane, aile sağlığı merkezi, emniyet, itfaiye, nüfus, noter, PTT, vergi dairesi, SGK, kütüphane, okul.',
  [('kategori',True,'Listeden seçin.'),('ad',True,'Kurumun resmî adı.')]+LOC+CONTACT+
  [('calisma_hafta_ici',False,'Örnek: 08:00-17:00'),('calisma_hafta_sonu',False,'Örnek: Kapalı'),('aciklama',False,'Kısa açıklama (en çok 300 karakter).')]+GEO+COMMON_TAIL,
  {'kategori':['valilik','kaymakamlik','il-mudurlugu','hastane-devlet','hastane-ozel','aile-sagligi-merkezi','emniyet','jandarma','itfaiye','nufus-mudurlugu','goc-idaresi','noter','ptt','vergi-dairesi','sgk','adliye','kutuphane','okul','halk-egitim','diger']},
  ['hastane-devlet','Bursa Şehir Hastanesi','Bursa','Nilüfer','Doğanköy','Doğanköy Mah. ...','','https://bursasehir.saglik.gov.tr','7/24','7/24','Acil servis 7/24 açık.','','','https://bursasehir.saglik.gov.tr','27.10.2026','ÖRNEK SATIR: bilgileri kontrol edilmedi, silin']),
 ('Müzeler','Devlet, belediye ve özel müzeler.',
  [('tur',True,'Listeden seçin.'),('ad',True,'Müzenin adı.')]+LOC+CONTACT+
  [('acilis_saatleri',False,'Örnek: 09:00-18:00 (yaz), 09:00-17:00 (kış)'),('kapali_gunler',False,'Örnek: Pazartesi'),
   ('ucret_durumu',True,'Ücretsiz / Ücretli'),('tam_ucret_tl',False,'Sayı (örnek: 150)'),('indirimli_ucret_tl',False,'Sayı'),
   ('muzekart_gecerli',False,'Evet / Hayır'),('aciklama',False,'Kısa tanıtım (en çok 500 karakter).')]+GEO+COMMON_TAIL,
  {'tur':['tarih','arkeoloji','sanat','sanayi','kultur','etnografya','bilim','kent','diger'],'ucret_durumu':['Ücretsiz','Ücretli'],'muzekart_gecerli':['Evet','Hayır']},
  ['kent','Bursa Kent Müzesi','Bursa','Osmangazi','','Reyhan Mah. ...','','https://www.bursa.bel.tr','10:00-18:00','Pazartesi','Ücretli','20','10','Hayır','Bursa\'nın kent tarihini anlatan müze.','','','https://www.bursa.bel.tr','27.10.2026','ÖRNEK SATIR: bilgileri kontrol edilmedi, silin']),
 ('Otoparklar','Belediye ve özel otoparklar (AVM dahil).',
  [('tip',True,'Listeden seçin.'),('ad',True,'Otoparkın adı.'),('isletici',False,'Belediye / Özel')]+LOC+CONTACT+
  [('kapasite',False,'Araç sayısı (sayı)'),('ilk_saat_tl',False,'Sayı'),('ek_saat_tl',False,'Sayı'),('gunluk_en_fazla_tl',False,'Sayı'),('aylik_abonelik_tl',False,'Sayı'),
   ('calisma_hafta_ici',False,'Örnek: 07:00-23:00'),('calisma_hafta_sonu',False,'Örnek: 08:00-24:00'),
   ('odeme_yontemleri',False,'Virgülle: Nakit, Kredi Kartı, Temassız, Mobil'),('engelli_ucretsiz',False,'Evet / Hayır'),('ozellikler',False,'Virgülle: Kapalı, Güvenlik kamerası, Elektrikli araç şarjı ...')]+GEO+COMMON_TAIL,
  {'tip':['kapali','acik','katli','yeralti','avm'],'isletici':['Belediye','Özel'],'engelli_ucretsiz':['Evet','Hayır']},
  ['katli','Örnek Katlı Otopark','Belediye','Bursa','Osmangazi','','... Cad. No:1','','','400','30','15','150','1500','07:00-23:00','08:00-24:00','Nakit, Kredi Kartı','Evet','Güvenlik kamerası','','','https://...','27.10.2026','ÖRNEK SATIR: bilgileri kontrol edilmedi, silin']),
 ('Parklar','Parklar, millet bahçeleri, kent ormanları, mesire alanları.',
  [('tur',True,'Listeden seçin.'),('ad',True,'Parkın adı.')]+LOC+
  [('acilis_saatleri',False,'Örnek: 7/24 ya da 06:00-24:00'),('olanaklar',False,'Virgülle: Çocuk oyun alanı, Spor aletleri, Yürüyüş yolu, Tuvalet, Otopark, Köpek alanı, Mescit, Kafe'),
   ('engelli_erisimi',False,'Evet / Hayır'),('aciklama',False,'Kısa tanıtım.')]+GEO+COMMON_TAIL,
  {'tur':['park','millet-bahcesi','kent-ormani','mesire-alani','botanik-bahcesi'],'engelli_erisimi':['Evet','Hayır']},
  ['park','Örnek Parkı','Bursa','Nilüfer','Ataevler','','7/24','Çocuk oyun alanı, Yürüyüş yolu','Evet','','','','https://...','27.10.2026','ÖRNEK SATIR: bilgileri kontrol edilmedi, silin']),
 ('Pazarlar','Semt pazarları, köylü/organik pazarlar, bit pazarları.',
  [('tur',True,'Listeden seçin.'),('ad',True,'Pazarın adı.')]+LOC+
  [('gunler',True,'Virgülle: Pazartesi, Salı, Çarşamba, Perşembe, Cuma, Cumartesi, Pazar'),('saatler',False,'Örnek: 08:00-18:00'),
   ('kapali_alan',False,'Evet (kapalı pazar yeri) / Hayır'),('urunler',False,'Virgülle: Sebze-meyve, Giyim, Organik, Köy ürünleri ...')]+GEO+COMMON_TAIL,
  {'tur':['semt-pazari','koylu-pazari','organik-pazar','bit-pazari','balik-pazari'],'kapali_alan':['Evet','Hayır']},
  ['semt-pazari','Örnek Salı Pazarı','Bursa','Osmangazi','','','Salı','08:00-18:00','Hayır','Sebze-meyve, Giyim','','','https://...','27.10.2026','ÖRNEK SATIR: bilgileri kontrol edilmedi, silin']),
 ('Ulaşım','Raylı sistem, otobüs hatları, teleferik, vapur, dolmuş, paylaşımlı bisiklet, ulaşım kartı.',
  [('tur',True,'Listeden seçin.'),('il',True,'İl adı.'),('hat_kodu',False,'Örnek: M1, B/43'),('ad',True,'Hat ya da hizmet adı (örnek: BursaRay M1).'),
   ('baslangic',False,'İlk durak'),('bitis',False,'Son durak'),('durak_sayisi',False,'Sayı'),('ilk_sefer',False,'Örnek: 06:00'),('son_sefer',False,'Örnek: 00:00'),
   ('sefer_sikligi',False,'Örnek: 5-10 dk'),('tam_ucret_tl',False,'Sayı'),('ogrenci_ucret_tl',False,'Sayı'),('kart_adi',False,'Örnek: BursaKart'),
   ('isletici',False,'Örnek: BURULAŞ'),('web',False,'Resmî web adresi.'),('aciklama',False,'Kısa bilgi.')]+COMMON_TAIL,
  {'tur':['metro','tramvay','banliyo','otobus-hatti','teleferik','vapur','dolmus','bisiklet','ulasim-karti','havalimani-servisi']},
  ['metro','Bursa','M1','BursaRay M1','Emek','Arabayatağı','','06:00','00:00','5-10 dk','','','BursaKart','BURULAŞ','https://www.burulas.com.tr','','https://www.burulas.com.tr','27.10.2026','ÖRNEK SATIR: bilgileri kontrol edilmedi, silin']),
 ('Toplanma Alanları','Afet ve acil durum toplanma alanları (AFAD).',
  [('ad',True,'Alanın adı.')]+LOC+[('kapasite_kisi',False,'Sayı')]+GEO+COMMON_TAIL,
  {},
  ['Örnek Okulu Bahçesi','Bursa','Nilüfer','Ataevler','','','','','https://www.turkiye.gov.tr/afet-ve-acil-durum-yonetimi-acil-toplanma-alani-sorgulama','27.10.2026','ÖRNEK SATIR: bilgileri kontrol edilmedi, silin']),
 ('Muhtarlıklar','Mahalle muhtarlıklarının resmî iletişim bilgileri.',
  [('il',True,'İl adı.'),('ilce',True,'İlçe adı.'),('mahalle',True,'Mahalle adı.'),('muhtar_adi',False,'Muhtarın adı soyadı (kamu görevlisi; yalnızca resmî kaynakta yayımlanmışsa).'),
   ('telefon',False,'Yalnızca muhtarlığın resmî numarası. Kişisel cep numarası yazılmaz.'),('adres',False,'Muhtarlık adresi.'),('calisma_saatleri',False,'Örnek: Hafta içi 09:00-17:00')]+GEO+COMMON_TAIL,
  {},
  ['Bursa','Nilüfer','Ataevler','','','','Hafta içi 09:00-17:00','','','https://...','27.10.2026','ÖRNEK SATIR: bilgileri kontrol edilmedi, silin']),
 ('Belediye Hizmetleri','Belediyenin verdiği hizmetler, gereken evrak ve başvuru bağlantısı (başvuru belediyenin kendi sitesinde yapılır).',
  [('il',True,'İl adı.'),('belediye',True,'Belediyenin tam adı (örnek: Nilüfer Belediyesi, Bursa Büyükşehir Belediyesi).'),('kategori',True,'Listeden seçin.'),
   ('hizmet_adi',True,'Örnek: Ruhsat başvurusu'),('aciklama',False,'Kısa açıklama.'),('gerekli_evrak',False,'Virgülle ayırın.'),('ucret',False,'Örnek: Ücretsiz ya da tutar'),
   ('basvuru_baglantisi',False,'Belediyenin başvuru sayfası.'),('telefon',False,'Resmî numara (örnek: 153).')]+COMMON_TAIL,
  {'kategori':['imar-ruhsat','vergi-odeme','su-altyapi','temizlik','zabita','sosyal-yardim','kultur-sanat','spor','egitim-kurs','cenaze','veteriner','diger']},
  ['Bursa','Nilüfer Belediyesi','imar-ruhsat','Yapı ruhsatı başvurusu','','Tapu, proje, ...','','https://...','153','https://...','27.10.2026','ÖRNEK SATIR: bilgileri kontrol edilmedi, silin']),
 ('Nasıl Yapılır','Şehir hayatı rehberi: resmî işlemler, ulaşım kartı, tasarruf ve pratik bilgiler. İl boş bırakılırsa ülke geneli sayılır.',
  [('il',False,'Boş = tüm Türkiye.'),('kategori',True,'Listeden seçin.'),('baslik',True,'Örnek: BursaKart nasıl alınır?'),('kisa_metin',True,'En çok 500 karakter.'),
   ('adimlar',False,'Adımları " | " ile ayırın.'),('gerekli_belgeler',False,'Virgülle ayırın.'),('ilgili_baglanti',False,'Resmî bağlantı.')]+COMMON_TAIL,
  {'kategori':['resmi-islem','ulasim','tasarruf','zaman','pratik','saglik','egitim']},
  ['Bursa','ulasim','BursaKart nasıl alınır?','BursaKart, şehir içi toplu taşımada kullanılan karttır.','Satış noktasına gidin | Kimlik gösterin | Kartı yükleyin','Kimlik (indirimli kart için)','https://...','https://...','27.10.2026','ÖRNEK SATIR: bilgileri kontrol edilmedi, silin']),
]

# İlçeler sekmesi (doğrulama listeleri için)
wl=wb.create_sheet('İlçeler')
wl.append(['il_kodu','il','ilce'])
for c in wl[1]: c.font=H; c.fill=HF
for d in districts:
    wl.append([d['provinceCode'],pname[d['provinceCode']],d['name']])
for c in ('A','B','C'): wl.column_dimensions[c].width=18 if c!='A' else 9
wl.freeze_panes='A2'
il_start=2
# İller listesi C sütunu dışında ayrı E sütununa
wl['E1']='iller'; wl['E1'].font=H; wl['E1'].fill=HF
for i,p in enumerate(provinces):
    wl.cell(row=2+i,column=5,value=p['name']).font=N
wl.column_dimensions['E'].width=18
for row in wl.iter_rows(min_row=2,max_col=3):
    for c in row: c.font=N

MAXR=5000
for title,desc,cols,choices,example in SHEETS:
    s=wb.create_sheet(title)
    s.append([c[0] for c in cols])
    for i,(name,req,hint) in enumerate(cols,1):
        cell=s.cell(row=1,column=i); cell.font=H; cell.fill=HF; cell.alignment=Alignment(vertical='center')
        cell.comment=Comment(('ZORUNLU. ' if req else 'İsteğe bağlı. ')+hint,'CityJoy')
        s.column_dimensions[get_column_letter(i)].width=max(14,min(40,len(name)+6))
    s.append(example)
    for c in s[2]: c.font=EX
    # zorunlu sütunların veri alanını hafif sarıyla işaretle (ilk 200 satır)
    for i,(name,req,hint) in enumerate(cols,1):
        col=get_column_letter(i)
        if req:
            for r in range(3,203): s[f'{col}{r}'].fill=REQ
        if name in choices:
            dv=DataValidation(type='list',formula1='"'+','.join(choices[name])+'"',allow_blank=True,showErrorMessage=True,
                              errorTitle='Listeden seçin',error='Bu sütunda yalnızca listedeki değerler kullanılır.')
            s.add_data_validation(dv); dv.add(f'{col}3:{col}{MAXR}')
        if name=='il':
            dv=DataValidation(type='list',formula1="'İlçeler'!$E$2:$E$82",allow_blank=True,showErrorMessage=True,
                              errorTitle='İl adı',error='İl adını listedeki yazımla seçin.')
            s.add_data_validation(dv); dv.add(f'{col}3:{col}{MAXR}')
    s.freeze_panes='A2'
    s.sheet_properties.tabColor='C9A227'

# Okuyun sekmesi
ws.column_dimensions['A'].width=28; ws.column_dimensions['B'].width=110
rows=[
 ('CityJoy · Benim Şehrim sabit bilgi şablonu',None),
 ('Sürüm','1 · 27 Ekim 2026'),
 ('',None),
 ('Nasıl doldurulur',None),
 ('1','Her sekme bir bölümdür. Her satır tek bir yerdir (bir müze, bir otopark, bir pazar...).'),
 ('2','Sütun adlarını ve sırasını değiştirmeyin, sekme adlarını değiştirmeyin. Fazla sütun eklemeyin.'),
 ('3','Her sekmenin 2. satırı gri ÖRNEK satırdır: biçimi gösterir, bilgileri doğrulanmamıştır. Doldururken silin ya da üzerine yazın.'),
 ('4','Sarı alanlar zorunludur. Başlık hücresinin üstüne gelince o sütunun açıklaması görünür.'),
 ('5','İl ve ilçe adlarını "İlçeler" sekmesindeki yazımla yazın (örnek: "Nilüfer", "Merkez"). İl sütununda açılır liste vardır.'),
 ('6','Listeden seçilen sütunlarda (kategori, tür, Evet/Hayır) yalnızca listedeki değerler kullanılır.'),
 ('7','Her satırda KAYNAK zorunludur: bilginin alındığı resmî sayfanın adresi ve bakıldığı gün. Kaynağı olmayan satır yüklenmez.'),
 ('8','Bilinmeyen bilgi boş bırakılır. Tahmin edilmez, uydurulmaz. Özellikle telefon, adres, fiyat ve harita konumu.'),
 ('9','Kişisel veri yazılmaz: kişisel cep telefonu, e-posta, kimlik no. Yalnızca kurumların resmî iletişim bilgileri.'),
 ('10','Siyasi parti adı, kişi övgüsü, reklam yazılmaz (tarafsızlık ilkesi).'),
 ('11','Metinler Türkçe yazılır; diğer dillere çeviri sonra yapılır.'),
 ('12','Bir il için tek dosya ya da birçok il bir arada olabilir. Dosya adı örneği: benim-sehrim-bursa.xlsx'),
 ('',None),
 ('Yüklendikten sonra',None),
 ('•','Dosya panelden yüklenir; hatalı satırlar yüklemeden önce listelenir.'),
 ('•','Yüklenen bilgiler uygulamada "doğrulanmadı" işaretiyle ve kaynağıyla görünür; il yöneticisi, belediye ya da CityJoy ekibi doğruladıkça işaret kalkar.'),
 ('•','Kafe-restoran ve işletmeler bu şablona girmez: İşletmem bölümünden gelir. Nöbetçi eczane ve etkinlikler de ayrıdır.'),
 ('',None),
 ('Başka bir yapay zekâya verilecek metin',None),
 ('Metin',"Ekteki Excel şablonunu [İL ADI] ili için doldur. Kurallar: (1) Sütun adlarını, sırasını ve sekme adlarını değiştirme. (2) Her satır tek bir yer olsun. (3) Yalnızca gerçekten var olan ve resmî ya da güvenilir bir kaynakta yayımlanmış bilgileri yaz; her satıra 'kaynak' sütununa o bilginin bulunduğu sayfanın tam adresini, 'kaynak_tarihi' sütununa bugünün tarihini yaz. (4) Emin olmadığın hiçbir bilgiyi tahmin etme veya uydurma; telefon, adres, fiyat, saat ve harita konumunu bilmiyorsan boş bırak. (5) Kişisel telefon, kişisel e-posta yazma; yalnızca kurumların resmî iletişim bilgileri. (6) Liste sütunlarında (kategori, tur, Evet/Hayır) yalnızca açıklamadaki değerleri kullan. (7) İl ve ilçe adlarını 'İlçeler' sekmesindeki yazımla yaz. (8) Siyasi parti adı, övgü ve reklam yazma. (9) Gri örnek satırları sil. Sonuç olarak aynı biçimde .xlsx dosyası ver."),
 ('Uyarı','Yapay zekâ araçları bazen var olmayan yerler, yanlış telefon ya da adres üretebilir. Bu yüzden kaynak sütunu zorunlu; kaynağı açılmayan ya da tutmayan satırlar yüklenmeden ayıklanır.'),
 ('',None),
 ('Sekmeler',None),
]+[(t,d) for t,d,_,_,_ in SHEETS]+[('İlçeler','81 il ve 973 ilçenin doğru yazımı (PTT listesinden). Değiştirmeyin.')]
for a,b in rows:
    ws.append([a,b])
    r=ws.max_row
    ws.cell(r,1).font=N; ws.cell(r,2).font=N
    ws.cell(r,2).alignment=Alignment(wrap_text=True,vertical='top'); ws.cell(r,1).alignment=Alignment(vertical='top')
    if b is None and a:
        ws.cell(r,1).font=Font(name=F,bold=True,size=14 if r==1 else 12,color='0B3B47')
ws.sheet_properties.tabColor='0B3B47'
# İlçeler sekmesini sona taşı
wb.move_sheet('İlçeler', offset=len(wb.sheetnames))
out=f'{ROOT}/docs/veri-sablonlari/benim-sehrim-sablon.xlsx'
wb.save(out); print(out, wb.sheetnames)
